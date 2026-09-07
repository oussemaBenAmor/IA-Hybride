"""
Tests unitaires du service d'exécution de scénarios (app/services/scenario_runner.py).

Ces tests ciblent les fonctions PURES du module — lecture et normalisation des
fichiers de scénarios (CSV et Excel), interprétation des colonnes, génération du
modèle — sans aucun appel réseau ni base de données. Ils vérifient donc la partie
« logique » du runner, indépendamment de l'exécution réelle contre l'API /chat.

À placer dans : fastapi-service/tests/test_scenario_runner.py
Lancement :     pytest tests/test_scenario_runner.py
"""
import io

import pytest
from openpyxl import load_workbook

from app.services.scenario_runner import (
    _parse_optional_bool,
    _to_scenario,
    parse_scenarios,
    build_template_xlsx,
)


# ══════════════════════════════════════════════════════════════════════════════
# _parse_optional_bool : interprétation de la colonne expected_params_needed
# ══════════════════════════════════════════════════════════════════════════════
class TestParseOptionalBool:

    @pytest.mark.parametrize("value", ["true", "TRUE", "True", "1", "oui", "yes", "vrai", "x", "  true  "])
    def test_truthy_values_return_true(self, value):
        assert _parse_optional_bool(value) is True

    @pytest.mark.parametrize("value", ["false", "FALSE", "0", "non", "no", "faux"])
    def test_falsy_values_return_false(self, value):
        assert _parse_optional_bool(value) is False

    def test_empty_string_returns_none(self):
        # Vide = critère non renseigné → ne doit PAS être vérifié.
        assert _parse_optional_bool("") is None
        assert _parse_optional_bool("   ") is None

    def test_none_returns_none(self):
        assert _parse_optional_bool(None) is None

    def test_unknown_value_returns_none(self):
        # Une valeur non reconnue est traitée comme « non renseigné », pas comme une erreur.
        assert _parse_optional_bool("peut-etre") is None


# ══════════════════════════════════════════════════════════════════════════════
# _to_scenario : normalisation d'une ligne brute en scénario
# ══════════════════════════════════════════════════════════════════════════════
class TestToScenario:

    def test_minimal_valid_row(self):
        rec = {"test_id": "T1", "messages": "bonjour"}
        s = _to_scenario(rec)
        assert s is not None
        assert s["test_id"] == "T1"
        assert s["messages"] == "bonjour"
        # Les critères non fournis restent None.
        assert s["expected_case"] is None
        assert s["expected_fallback"] is None
        assert s["expected_odm_decision"] is None
        assert s["expected_params_needed"] is None

    def test_full_row_is_normalised(self):
        rec = {
            "test_id": "T2", "description": "desc", "category": "nominal",
            "messages": "je veux un credit | 200000 euros",
            "expected_case": "credit_immobilier",
            "expected_fallback": "",
            "expected_odm_decision": "eligible",     # doit être mis en MAJUSCULES
            "expected_params_needed": "true",
        }
        s = _to_scenario(rec)
        assert s["category"] == "nominal"
        assert s["expected_case"] == "credit_immobilier"
        assert s["expected_fallback"] is None        # chaîne vide → None
        assert s["expected_odm_decision"] == "ELIGIBLE"
        assert s["expected_params_needed"] is True

    def test_row_without_messages_is_ignored(self):
        # Sans messages, le scénario n'est pas exploitable.
        assert _to_scenario({"test_id": "T3", "messages": ""}) is None
        assert _to_scenario({"test_id": "T3"}) is None

    def test_completely_empty_row_is_ignored(self):
        assert _to_scenario({"test_id": "", "messages": ""}) is None

    def test_whitespace_is_stripped(self):
        s = _to_scenario({"test_id": "  T4  ", "messages": "  bonjour  "})
        assert s["test_id"] == "T4"
        assert s["messages"] == "bonjour"


# ══════════════════════════════════════════════════════════════════════════════
# parse_scenarios : lecture d'un fichier CSV
# ══════════════════════════════════════════════════════════════════════════════
class TestParseCsv:

    def test_simple_csv(self):
        csv = b"test_id,messages\nT1,bonjour\nT2,au revoir\n"
        scenarios = parse_scenarios(csv, "cases.csv")
        assert len(scenarios) == 2
        assert scenarios[0]["test_id"] == "T1"
        assert scenarios[1]["messages"] == "au revoir"

    def test_csv_column_order_is_irrelevant(self):
        # Les colonnes sont lues par NOM : l'ordre ne doit rien changer.
        csv = (b"messages,expected_case,test_id\n"
               b"je veux un pret,credit_consommation,T9\n")
        scenarios = parse_scenarios(csv, "cases.csv")
        assert scenarios[0]["test_id"] == "T9"
        assert scenarios[0]["expected_case"] == "credit_consommation"

    def test_csv_with_semicolon_separator(self):
        # Le séparateur ';' (Excel FR) doit être détecté automatiquement.
        csv = b"test_id;messages\nT1;bonjour\n"
        scenarios = parse_scenarios(csv, "cases.csv")
        assert len(scenarios) == 1
        assert scenarios[0]["test_id"] == "T1"

    def test_csv_with_utf8_bom(self):
        # Un CSV exporté par Excel commence souvent par un BOM UTF-8.
        csv = "\ufefftest_id,messages\nT1,héberger\n".encode("utf-8")
        scenarios = parse_scenarios(csv, "cases.csv")
        assert scenarios[0]["test_id"] == "T1"
        assert "héberger" in scenarios[0]["messages"]

    def test_csv_multiturn_messages_preserved(self):
        # Le séparateur de tours '|' est conservé tel quel (découpé plus tard à l'exécution).
        csv = b"test_id,messages\nT1,je veux un credit|200000 euros\n"
        scenarios = parse_scenarios(csv, "cases.csv")
        assert scenarios[0]["messages"] == "je veux un credit|200000 euros"

    def test_csv_missing_required_column_raises(self):
        # Sans la colonne obligatoire 'messages', on doit lever une erreur explicite.
        csv = b"test_id,category\nT1,nominal\n"
        with pytest.raises(ValueError, match="messages"):
            parse_scenarios(csv, "cases.csv")

    def test_empty_csv_raises(self):
        with pytest.raises(ValueError):
            parse_scenarios(b"test_id,messages\n", "cases.csv")

    def test_blank_rows_are_skipped(self):
        csv = b"test_id,messages\nT1,bonjour\n,\nT2,merci\n"
        scenarios = parse_scenarios(csv, "cases.csv")
        # La ligne vide du milieu est ignorée.
        assert [s["test_id"] for s in scenarios] == ["T1", "T2"]


# ══════════════════════════════════════════════════════════════════════════════
# parse_scenarios : lecture d'un fichier Excel (.xlsx)
# ══════════════════════════════════════════════════════════════════════════════
class TestParseXlsx:

    def _make_xlsx(self, rows: list[list]) -> bytes:
        """Construit un .xlsx en mémoire à partir d'une liste de lignes."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_simple_xlsx(self):
        data = self._make_xlsx([
            ["test_id", "messages", "expected_case"],
            ["T1", "je veux un credit immobilier", "credit_immobilier"],
        ])
        scenarios = parse_scenarios(data, "cases.xlsx")
        assert len(scenarios) == 1
        assert scenarios[0]["expected_case"] == "credit_immobilier"

    def test_xlsx_headers_case_insensitive(self):
        # Les en-têtes sont normalisés en minuscules.
        data = self._make_xlsx([
            ["Test_ID", "Messages"],
            ["T1", "bonjour"],
        ])
        scenarios = parse_scenarios(data, "cases.xlsx")
        assert scenarios[0]["test_id"] == "T1"

    def test_xlsx_missing_required_column_raises(self):
        data = self._make_xlsx([["test_id", "category"], ["T1", "nominal"]])
        with pytest.raises(ValueError, match="messages"):
            parse_scenarios(data, "cases.xlsx")


# ══════════════════════════════════════════════════════════════════════════════
# parse_scenarios : détection du format et robustesse
# ══════════════════════════════════════════════════════════════════════════════
class TestFormatDetection:

    def test_unknown_extension_falls_back_to_parsing(self):
        # Sans extension claire, le contenu CSV doit tout de même être lu.
        csv = b"test_id,messages\nT1,bonjour\n"
        scenarios = parse_scenarios(csv, "cases.data")
        assert scenarios[0]["test_id"] == "T1"

    def test_no_exploitable_scenario_raises(self):
        # En-têtes présents mais aucune ligne exploitable.
        csv = b"test_id,messages\n,\n,\n"
        with pytest.raises(ValueError):
            parse_scenarios(csv, "cases.csv")


# ══════════════════════════════════════════════════════════════════════════════
# build_template_xlsx : le modèle vierge est un Excel valide et exploitable
# ══════════════════════════════════════════════════════════════════════════════
class TestBuildTemplate:

    def test_template_is_valid_xlsx(self):
        data = build_template_xlsx()
        assert isinstance(data, bytes) and len(data) > 0
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        # Les colonnes obligatoires et la nouvelle colonne doivent être présentes.
        assert "test_id" in headers
        assert "messages" in headers
        assert "expected_params_needed" in headers

    def test_template_can_be_reparsed(self):
        # Le modèle généré doit lui-même être relisible par parse_scenarios.
        data = build_template_xlsx()
        scenarios = parse_scenarios(data, "modele.xlsx")
        assert len(scenarios) >= 1
        assert all("test_id" in s and "messages" in s for s in scenarios)