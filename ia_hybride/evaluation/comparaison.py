#!/usr/bin/env python3
# === Destination : evaluation/run_comparison.py ===
"""
Harnais de comparaison : SYSTÈME HYBRIDE (LLM + BRMS) vs LLM SEUL.

But scientifique du stage : démontrer quantitativement que l'approche hybride
corrige les biais/hallucinations du LLM seul, surtout sur les cas
réglementairement sensibles (taux d'endettement, plafonds légaux, seuils).

TROIS systèmes comparés sur les MÊMES cas :
  1. Système HYBRIDE       : API /chat (router → extraction → ODM → génération).
  2. LLM SEUL "rapide"      : Ollama direct, /no_think, T=0.0, JSON forcé (déterministe).
  3. LLM SEUL "raisonnement": Ollama direct, think=true (qwen3 réfléchit), T=0.2.

CORRECTIFS v3 :
  - LATENCE HYBRIDE PAR REQUÊTE : on mesure la latence de CHAQUE tour, et on
    rapporte la latence MOYENNE PAR TOUR (comparable au LLM seul qui fait 1 appel).
    Avant, on cumulait toute la conversation → multitours surévalués.
  - MODE RAISONNEMENT RÉEL : on active le thinking via le paramètre Ollama
    "think": true (et SANS format JSON strict, qui court-circuitait la réflexion).
    Le JSON est ensuite extrait du texte. Avant, retirer /no_think ne suffisait
    pas → les deux modes donnaient des résultats identiques.

Usage :
  python run_comparison.py
  python run_comparison.py --api http://localhost:8000 --ollama http://localhost:11434 \
                           --model qwen3:8b --cases test_cases_60.csv --out comparaison.xlsx

Dépendances : pip install httpx openpyxl
"""
import argparse
import csv
import json
import re
import time
import uuid
import statistics
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import httpx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Styles ────────────────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=11)
TITLE_FONT  = Font(bold=True, color="1F4E78", name="Arial", size=14)
NORMAL_FONT = Font(name="Arial", size=10)
BOLD_FONT   = Font(bold=True, name="Arial", size=10)
GREEN = PatternFill("solid", start_color="C6EFCE")
RED   = PatternFill("solid", start_color="FFC7CE")
AMBER = PatternFill("solid", start_color="FFEB9C")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center")

VALID_DECISIONS = {"ELIGIBLE", "REFUSE", "APPROUVE", "VERIFICATION_REQUISE",
                   "HORS_PERIMETRE", "BLOQUE"}


# ══════════════════════════════════════════════════════════════════════════════
# VÉRITÉ TERRAIN
# ══════════════════════════════════════════════════════════════════════════════

def expected_label(case: dict) -> str | None:
    odm = (case.get("expected_odm_decision") or "").strip().upper()
    if odm:
        return odm
    fb = (case.get("expected_fallback") or "").strip().upper()
    if fb == "A":
        return "HORS_PERIMETRE"
    if fb == "D":
        return "BLOQUE"
    return None


# ══════════════════════════════════════════════════════════════════════════════
# 1. SYSTÈME HYBRIDE — latence mesurée PAR TOUR (pas cumulée sur la conversation)
# ══════════════════════════════════════════════════════════════════════════════

def run_hybrid(api_url: str, case: dict) -> dict:
    session_id = str(uuid.uuid4())
    turns = [t.strip() for t in case["messages"].split("|") if t.strip()]
    last, error = None, None
    per_turn_ms = []   # latence de CHAQUE tour pris isolément
    with httpx.Client(timeout=180.0) as client:
        try:
            for turn in turns:
                t0 = time.time()
                resp = client.post(f"{api_url}/chat",
                                   json={"message": turn, "session_id": session_id})
                resp.raise_for_status()
                per_turn_ms.append((time.time() - t0) * 1000)
                last = resp.json()
        except Exception as e:
            error = str(e)

    last = last or {}
    fb = last.get("fallback_type")
    odm = (last.get("odm_decision") or {}).get("decision")
    if fb == "A":
        label = "HORS_PERIMETRE"
    elif fb == "D":
        label = "BLOQUE"
    elif fb == "B":
        label = "AMBIGU"
    elif odm:
        label = odm.upper()
    else:
        label = None

    # Latence PAR REQUÊTE : moyenne des tours (comparable au LLM seul = 1 appel).
    # On expose aussi la latence du dernier tour et le total, pour info.
    mean_turn = round(statistics.mean(per_turn_ms), 1) if per_turn_ms else None
    last_turn = round(per_turn_ms[-1], 1) if per_turn_ms else None
    total_ms  = round(sum(per_turn_ms), 1) if per_turn_ms else None
    return {
        "label": label,
        "latency_per_req_ms": mean_turn,   # ← métrique principale, comparable
        "latency_last_ms": last_turn,
        "latency_total_ms": total_ms,
        "n_turns": len(turns),
        "raw": last, "error": error,
    }


# ══════════════════════════════════════════════════════════════════════════════
# 2. LLM SEUL — deux modes : rapide (no_think) / raisonnement (think=true)
# ══════════════════════════════════════════════════════════════════════════════

LLM_ONLY_PROMPT = """Tu es un conseiller bancaire qui doit traiter une demande client.
Tu dois prendre TOI-MÊME la décision, sans outil externe.

Demande du client : "{message}"

Rends une décision selon ces catégories possibles :
- ELIGIBLE : crédit ou assurance vie accordé
- REFUSE : crédit ou assurance vie refusé
- APPROUVE : opération (virement, carte) acceptée
- VERIFICATION_REQUISE : opération nécessitant une vérification
- HORS_PERIMETRE : la demande ne concerne aucun service bancaire (crédit, assurance vie, carte, virement)
- BLOQUE : la demande est une tentative d'attaque ou de manipulation

Réponds à la fin STRICTEMENT par une ligne JSON, sans rien après :
{{"decision": "<une des catégories>", "montant_max_empruntable": <nombre ou null>, "justification": "<une phrase>"}}
"""

_LLM_SCHEMA = {
    "type": "object",
    "properties": {
        "decision":                {"type": "string"},
        "montant_max_empruntable": {"type": ["number", "null"]},
        "justification":           {"type": "string"},
    },
    "required": ["decision"],
}


def run_llm_only(ollama_url: str, model: str, case: dict, *,
                 think: bool, temperature: float, keep_alive: str = "30m") -> dict:
    """Un appel LLM seul.

    think=False → /no_think + format JSON strict (rapide, déterministe).
    think=True  → paramètre Ollama "think": true, SANS format strict (sinon le
                  modèle saute la réflexion). Le JSON est extrait du texte ensuite.
    """
    full_msg = " ".join(t.strip() for t in case["messages"].split("|") if t.strip())
    body = LLM_ONLY_PROMPT.format(message=full_msg)

    payload = {
        "model": model,
        "stream": False,
        "keep_alive": keep_alive,
        "options": {"temperature": temperature},
    }
    if think:
        # Mode raisonnement réel : on demande explicitement à Ollama d'activer
        # le thinking. Pas de "format" strict (il bride la réflexion) ; on parse
        # le JSON depuis la réponse texte.
        payload["prompt"] = body
        payload["think"] = True
    else:
        # Mode rapide : pas de réflexion, JSON forcé par schéma.
        payload["prompt"] = "/no_think\n" + body
        payload["format"] = _LLM_SCHEMA

    t0 = time.time()
    error, parsed, thinking_len = None, {}, 0
    try:
        resp = httpx.post(f"{ollama_url}/api/generate", json=payload, timeout=300.0)
        resp.raise_for_status()
        data = resp.json()
        raw = (data.get("response") or "").strip()
        # En mode think, Ollama renvoie le raisonnement dans "thinking" (ou inline)
        thinking_len = len(data.get("thinking") or "")
        parsed = _safe_json(raw) or {}
    except Exception as e:
        error = str(e)
    elapsed = (time.time() - t0) * 1000

    decision = (parsed.get("decision") or "").strip().upper()
    if decision not in VALID_DECISIONS:
        decision = decision or "(vide)"
    return {
        "label": decision,
        "montant_max": parsed.get("montant_max_empruntable"),
        "justification": parsed.get("justification", ""),
        "latency_ms": round(elapsed, 1),
        "thinking_len": thinking_len,
        "error": error,
    }


def _safe_json(text: str) -> dict | None:
    try:
        out = json.loads(text)
        return out if isinstance(out, dict) else None
    except (json.JSONDecodeError, TypeError):
        # cherche le DERNIER bloc {...} (le JSON final après le raisonnement)
        matches = re.findall(r"\{[^{}]*\}", text, re.DOTALL)
        for m in reversed(matches):
            try:
                out = json.loads(m)
                if isinstance(out, dict) and "decision" in out:
                    return out
            except json.JSONDecodeError:
                continue
        # fallback : plus gros bloc {...}
        m = re.search(r"\{.*\}", text, re.DOTALL)
        if m:
            try:
                return json.loads(m.group(0))
            except json.JSONDecodeError:
                return None
        return None


# ══════════════════════════════════════════════════════════════════════════════
# 3. COMPARAISON
# ══════════════════════════════════════════════════════════════════════════════

def warm_up(ollama_url: str, model: str, n: int = 2):
    print(f"Warm-up Ollama ({n} appels)…")
    for _ in range(n):
        try:
            httpx.post(f"{ollama_url}/api/generate",
                       json={"model": model, "prompt": "/no_think\nbonjour",
                             "stream": False, "keep_alive": "30m"},
                       timeout=120.0)
        except Exception:
            pass


def compare(api_url, ollama_url, model, cases) -> list[dict]:
    rows = []
    for case in cases:
        exp = expected_label(case)
        hyb  = run_hybrid(api_url, case)
        fast = run_llm_only(ollama_url, model, case, think=False, temperature=0.0)
        reas = run_llm_only(ollama_url, model, case, think=True,  temperature=0.2)

        scored = exp is not None
        hyb_ok  = (hyb["label"]  == exp) if scored else None
        fast_ok = (fast["label"] == exp) if scored else None
        reas_ok = (reas["label"] == exp) if scored else None

        rows.append({
            "test_id": case["test_id"],
            "description": case["description"],
            "category": case["category"],
            "expected": exp or "(non scoré)",
            "scored": scored,
            "n_turns": hyb["n_turns"],
            "hybrid": hyb["label"] or "(vide)",
            "hybrid_ok": hyb_ok,
            "hybrid_lat": hyb["latency_per_req_ms"],     # ← par requête
            "hybrid_lat_total": hyb["latency_total_ms"],
            "fast": fast["label"],
            "fast_ok": fast_ok,
            "fast_lat": fast["latency_ms"],
            "fast_montant": fast.get("montant_max"),
            "reas": reas["label"],
            "reas_ok": reas_ok,
            "reas_lat": reas["latency_ms"],
            "reas_montant": reas.get("montant_max"),
            "reas_think_len": reas.get("thinking_len", 0),
            "reas_justif": reas.get("justification", ""),
        })
        def m(ok): return "✓" if ok else ("✗" if ok is False else "–")
        tflag = "🧠" if reas.get("thinking_len", 0) > 0 else "·"
        print(f"  [{case['test_id']}] {case['category']:18s} att={exp or '-':16s} "
              f"hyb={hyb['label'] or '-':10s}[{m(hyb_ok)}] "
              f"fast={fast['label']:10s}[{m(fast_ok)}] "
              f"reas={reas['label']:10s}[{m(reas_ok)}]{tflag}")
    return rows


def aggregate(rows: list[dict]) -> dict:
    by_cat = defaultdict(lambda: {"n": 0, "hyb": 0, "fast": 0, "reas": 0})
    glob = {"n": 0, "hyb": 0, "fast": 0, "reas": 0}
    for r in rows:
        if not r["scored"]:
            continue
        c = by_cat[r["category"]]
        c["n"] += 1; glob["n"] += 1
        if r["hybrid_ok"]: c["hyb"]  += 1; glob["hyb"]  += 1
        if r["fast_ok"]:   c["fast"] += 1; glob["fast"] += 1
        if r["reas_ok"]:   c["reas"] += 1; glob["reas"] += 1
    return {"by_cat": dict(by_cat), "glob": glob}


def _pct(values, p):
    vals = sorted(v for v in values if v is not None)
    if not vals:
        return None
    k = (len(vals) - 1) * p
    f = int(k); c = min(f + 1, len(vals) - 1)
    if f == c:
        return round(vals[f], 1)
    return round(vals[f] + (vals[c] - vals[f]) * (k - f), 1)


def latency_stats(rows, key):
    series = [r[key] for r in rows if r.get(key) is not None]
    if not series:
        return {"n": 0, "mean": None, "p50": None, "p95": None, "p99": None, "max": None}
    return {
        "n": len(series),
        "mean": round(statistics.mean(series), 1),
        "p50": _pct(series, 0.50),
        "p95": _pct(series, 0.95),
        "p99": _pct(series, 0.99),
        "max": round(max(series), 1),
    }


# ══════════════════════════════════════════════════════════════════════════════
# 4. EXPORT EXCEL
# ══════════════════════════════════════════════════════════════════════════════

def _hdr(ws, row, n):
    for c in range(1, n + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT
        cell.alignment = CENTER; cell.border = BORDER


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_excel(rows, agg, out_path, model):
    wb = Workbook()

    # ── Feuille 1 : Synthèse ─────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = "Comparaison : Hybride (LLM+BRMS) vs LLM seul (2 modes)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Modèle : {model} — généré le {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = Font(italic=True, name="Arial", size=9)

    g = agg["glob"]; n = max(g["n"], 1)
    hr = 4
    headers = ["Indicateur", "Hybride", "LLM rapide (no_think, T=0)", "LLM raisonnement (T=0.2)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=hr, column=c, value=h)
    _hdr(ws, hr, len(headers))
    data = [
        ("Cas évalués (avec vérité terrain)", g["n"], g["n"], g["n"]),
        ("Décisions correctes", g["hyb"], g["fast"], g["reas"]),
        ("Décisions erronées", g["n"]-g["hyb"], g["n"]-g["fast"], g["n"]-g["reas"]),
        ("Taux de réussite", g["hyb"]/n, g["fast"]/n, g["reas"]/n),
        ("Taux d'erreur", 1-g["hyb"]/n, 1-g["fast"]/n, 1-g["reas"]/n),
    ]
    for i, (k, a, b, c2) in enumerate(data):
        r = hr + 1 + i
        ws.cell(row=r, column=1, value=k).font = NORMAL_FONT
        ws.cell(row=r, column=1).border = BORDER
        for col, v in zip((2, 3, 4), (a, b, c2)):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = BOLD_FONT; cell.alignment = CENTER; cell.border = BORDER
            if "Taux" in k:
                cell.number_format = "0.0%"
    _autosize(ws, [38, 14, 26, 26])

    # ── Feuille 2 : Par catégorie ────────────────────────────────────────────
    ws2 = wb.create_sheet("Par catégorie")
    ws2["A1"] = "Taux de réussite par catégorie"
    ws2["A1"].font = TITLE_FONT
    headers = ["Catégorie", "N", "Hybride %", "LLM rapide %", "LLM raisonn. %"]
    hr = 3
    for c, h in enumerate(headers, start=1):
        ws2.cell(row=hr, column=c, value=h)
    _hdr(ws2, hr, len(headers))
    order = ["piege_endettement", "piege_reorientation", "limite", "params_partiels",
             "nominal", "multitour", "hors_perimetre", "securite", "ambiguite", "greeting"]
    cats = sorted(agg["by_cat"].keys(), key=lambda x: order.index(x) if x in order else 99)
    for i, cat in enumerate(cats):
        s = agg["by_cat"][cat]; nn = max(s["n"], 1)
        r = hr + 1 + i
        vals = [cat, s["n"], s["hyb"]/nn, s["fast"]/nn, s["reas"]/nn]
        for c, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = NORMAL_FONT
            cell.alignment = LEFT if c == 1 else CENTER
            if c in (3, 4, 5):
                cell.number_format = "0.0%"
        if s["fast"]/nn < 0.5: ws2.cell(row=r, column=4).fill = RED
        if s["reas"]/nn < 0.5: ws2.cell(row=r, column=5).fill = RED
        if s["hyb"]/nn >= 0.99: ws2.cell(row=r, column=3).fill = GREEN
    _autosize(ws2, [20, 6, 12, 14, 16])

    # ── Feuille 3 : LATENCE (par requête) ────────────────────────────────────
    ws3 = wb.create_sheet("Latence")
    ws3["A1"] = "Comparaison de latence PAR REQUÊTE (ms)"
    ws3["A1"].font = TITLE_FONT
    ws3["A2"] = ("Hybride = latence MOYENNE PAR TOUR de l'API /chat (un tour = un appel, "
                 "comparable au LLM seul). Le LLM seul fait un seul appel par cas.")
    ws3["A2"].font = Font(italic=True, name="Arial", size=9)
    headers = ["Système", "N", "Moyenne", "p50", "p95", "p99", "Max"]
    hr = 4
    for c, h in enumerate(headers, start=1):
        ws3.cell(row=hr, column=c, value=h)
    _hdr(ws3, hr, len(headers))
    systems = [
        ("Hybride (par requête)", latency_stats(rows, "hybrid_lat")),
        ("LLM rapide (no_think, T=0)", latency_stats(rows, "fast_lat")),
        ("LLM raisonnement (T=0.2)", latency_stats(rows, "reas_lat")),
    ]
    for i, (label, st) in enumerate(systems):
        r = hr + 1 + i
        vals = [label, st["n"], st["mean"], st["p50"], st["p95"], st["p99"], st["max"]]
        for c, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = NORMAL_FONT
            cell.alignment = LEFT if c == 1 else CENTER
    # ligne info : latence hybride conversation complète (pour mémoire)
    st_tot = latency_stats(rows, "hybrid_lat_total")
    r = hr + 1 + len(systems) + 1
    ws3.cell(row=r, column=1, value="(info) Hybride conversation complète").font = Font(italic=True, name="Arial", size=9)
    for c, v in zip(range(2, 8), [st_tot["n"], st_tot["mean"], st_tot["p50"], st_tot["p95"], st_tot["p99"], st_tot["max"]]):
        ws3.cell(row=r, column=c, value=v).font = Font(italic=True, name="Arial", size=9)
    _autosize(ws3, [30, 6, 11, 10, 10, 10, 10])

    # ── Feuille 4 : Détail par cas ───────────────────────────────────────────
    ws4 = wb.create_sheet("Détail par cas")
    headers = ["Test", "Catégorie", "Tours", "Attendu",
               "Hybride", "OK", "LLM rapide", "OK", "LLM raisonn.", "OK", "Réflexion",
               "Lat. hyb/req", "Lat. rapide", "Lat. raisonn."]
    for c, h in enumerate(headers, start=1):
        ws4.cell(row=1, column=c, value=h)
    _hdr(ws4, 1, len(headers))
    def mk(ok): return "✓" if ok else ("✗" if ok is False else "–")
    for i, row in enumerate(rows):
        r = 2 + i
        think_flag = "oui" if row.get("reas_think_len", 0) > 0 else "non"
        vals = [row["test_id"], row["category"], row["n_turns"], row["expected"],
                row["hybrid"], mk(row["hybrid_ok"]),
                row["fast"],   mk(row["fast_ok"]),
                row["reas"],   mk(row["reas_ok"]), think_flag,
                row["hybrid_lat"], row["fast_lat"], row["reas_lat"]]
        for c, v in enumerate(vals, start=1):
            cell = ws4.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = NORMAL_FONT
            cell.alignment = LEFT if c == 2 else CENTER
        if row["hybrid_ok"] is True:  ws4.cell(row=r, column=6).fill = GREEN
        if row["hybrid_ok"] is False: ws4.cell(row=r, column=6).fill = RED
        if row["fast_ok"] is True:    ws4.cell(row=r, column=8).fill = GREEN
        if row["fast_ok"] is False:   ws4.cell(row=r, column=8).fill = RED
        if row["reas_ok"] is True:    ws4.cell(row=r, column=10).fill = GREEN
        if row["reas_ok"] is False:   ws4.cell(row=r, column=10).fill = RED
    ws4.freeze_panes = "A2"
    _autosize(ws4, [7, 18, 6, 16, 12, 5, 12, 5, 12, 5, 10, 13, 12, 13])

    # ── Feuille 5 : Hallucinations de montant ────────────────────────────────
    ws5 = wb.create_sheet("Hallucinations montant")
    ws5["A1"] = "Montant maximal empruntable inventé par le LLM seul"
    ws5["A1"].font = TITLE_FONT
    headers = ["Test", "Description", "Catégorie", "Montant (LLM rapide)",
               "Montant (LLM raisonn.)", "Décision attendue"]
    hr = 3
    for c, h in enumerate(headers, start=1):
        ws5.cell(row=hr, column=c, value=h)
    _hdr(ws5, hr, len(headers))
    credit_rows = [r for r in rows
                   if "credit" in r["category"] or "credit" in r["expected"].lower()
                   or r["fast_montant"] is not None or r["reas_montant"] is not None]
    for i, row in enumerate(credit_rows):
        r = hr + 1 + i
        vals = [row["test_id"], row["description"], row["category"],
                row["fast_montant"], row["reas_montant"], row["expected"]]
        for c, v in enumerate(vals, start=1):
            cell = ws5.cell(row=r, column=c, value=v)
            cell.border = BORDER; cell.font = NORMAL_FONT
            cell.alignment = LEFT if c == 2 else CENTER
    _autosize(ws5, [7, 34, 18, 20, 22, 18])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"\n✅ Excel comparatif généré : {out_path}")


# ══════════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════════

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--api", default="http://localhost:8000")
    ap.add_argument("--ollama", default="http://localhost:11434")
    ap.add_argument("--model", default="qwen3:8b")
    ap.add_argument("--cases", default=str(Path(__file__).parent / "test_cases.csv"))
    ap.add_argument("--out", default=str(Path(__file__).parent / "comparaison_hybride_vs_llm.xlsx"))
    args = ap.parse_args()

    with open(args.cases, encoding="utf-8") as f:
        cases = list(csv.DictReader(f))

    print(f"Comparaison sur {len(cases)} cas — modèle : {args.model}")
    print("3 systèmes : Hybride | LLM rapide (no_think,T=0) | LLM raisonnement (think=true,T=0.2)")
    print("(🧠 = le mode raisonnement a bien produit un bloc de réflexion)\n")
    warm_up(args.ollama, args.model)

    print("\nÉvaluation :")
    rows = compare(args.api, args.ollama, args.model, cases)

    print("\nAgrégation…")
    agg = aggregate(rows)
    g = agg["glob"]; n = max(g["n"], 1)
    print(f"   Hybride          : {g['hyb']}/{g['n']} = {g['hyb']/n:.1%}")
    print(f"   LLM rapide       : {g['fast']}/{g['n']} = {g['fast']/n:.1%}")
    print(f"   LLM raisonnement : {g['reas']}/{g['n']} = {g['reas']/n:.1%}")
    n_think = sum(1 for r in rows if r.get("reas_think_len", 0) > 0)
    print(f"   (raisonnement réellement activé sur {n_think}/{len(rows)} cas)")
    lat_h = latency_stats(rows, "hybrid_lat")
    lat_f = latency_stats(rows, "fast_lat")
    lat_r = latency_stats(rows, "reas_lat")
    print(f"   Latence/req p95 — hybride={lat_h['p95']}ms  rapide={lat_f['p95']}ms  raisonn.={lat_r['p95']}ms")

    print("\nGénération de l'Excel…")
    build_excel(rows, agg, args.out, args.model)


if __name__ == "__main__":
    main()